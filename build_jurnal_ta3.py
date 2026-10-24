"""
build_jurnal_ta3.py - Laporan TA-03 format Jurnal Informatika COMPUTING (dua kolom) dari model/metrics.json
Jalankan setelah train.py:  python build_jurnal_ta3.py     (isi LINK_* di bawah lebih dulu)
"""
import json
import math
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ============================ KONFIGURASI (ISI) ============================
NAMA = "Muhammad Fadillah Ramadhan"
EMAIL = "fadilrmdn33@gmail.com"
NIM = "301240051"
LINK_GITHUB = "https://github.com/USERNAME/credit-risk-app"
LINK_APP = "https://NAMA-APP.onrender.com"
LINK_YOUTUBE = "https://youtu.be/XXXXXXXXXXX"
LINK_KAGGLE = "https://www.kaggle.com/datasets/uciml/default-of-credit-card-clients-dataset"
OUT = f"LAPORAN_TA03_PEMBELAJARAN_MESIN_{NAMA.upper().replace(' ', '_')}_{NIM}.docx"
# ===========================================================================

m = json.load(open(os.path.join("model", "metrics.json"), encoding="utf-8"))
IMG = os.path.join("static", "img")
META = json.load(open(os.path.join("model", "meta.json"), encoding="utf-8"))
LOGO = os.path.join("assets", "logo_jurnal.png")


def fx(x, n=2):
    """format angka gaya Indonesia (koma desimal)."""
    return f"{x:,.{n}f}".replace(",", "§").replace(".", ",").replace("§", ".")



doc = Document()
st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.space_before = Pt(0)
st.paragraph_format.line_spacing = 1.0


def fix_table(t, widths_cm):
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)


def set_cols(section, n, space_cm=0.8):
    sectPr = section._sectPr
    for e in sectPr.findall(qn("w:cols")):
        sectPr.remove(e)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space_cm * 567)))
    sectPr.append(cols)


def page_setup(s):
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)  # letter, seperti template
    s.left_margin = s.right_margin = Cm(1.9)
    s.top_margin = Cm(3.3)
    s.bottom_margin = Cm(2.0)
    s.header_distance = Cm(1.0)


sec1 = doc.sections[0]
page_setup(sec1)

# ------------------------------------------------------------------ HEADER
hdr = sec1.header
hdr.is_linked_to_previous = False
tbl = hdr.add_table(1, 2, width=Cm(17.7))
tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl.autofit = False
lc, rc = tbl.rows[0].cells
fix_table(tbl, [4.2, 13.5])
lc.paragraphs[0].add_run().add_picture(LOGO, width=Cm(3.1))
lines = [
    ("Jurnal Informatika – COMPUTING Volume xx Nomor xx, Bulan xxxxxx Tahun xxxx: xx – xx", False),
    ("ISSN: 2656 –386X", True),
    ("e-ISSN: 2722 –2888", True),
]
for i, (txt, bold) in enumerate(lines):
    p = rc.paragraphs[0] if i == 0 else rc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(txt)
    r.font.size = Pt(8)
    r.bold = bold
# garis bawah header
hp = hdr.paragraphs[0]
hp.paragraph_format.space_after = Pt(0)
pPr = hp._p.get_or_add_pPr()
bd = OxmlElement("w:pBdr")
bt = OxmlElement("w:bottom")
for k, v in (("val", "single"), ("sz", "10"), ("space", "1"), ("color", "000000")):
    bt.set(qn(f"w:{k}"), v)
bd.append(bt)
pPr.append(bd)
# paragraf kosong bawaan header dipindah ke bawah tabel
hdr._element.remove(hp._p)
hdr._element.append(hp._p)


# ------------------------------------------------------------------ HELPERS
def P(text="", bold=False, italic=False, align="j", size=None, color=None, after=None, keep=False):
    p = doc.add_paragraph()
    p.alignment = {"j": WD_ALIGN_PARAGRAPH.JUSTIFY, "c": WD_ALIGN_PARAGRAPH.CENTER,
                   "l": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if keep:
        p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return p


def H(text):
    p = P(text, bold=True, align="l", keep=True, after=4)
    p.paragraph_format.space_before = Pt(8)


def SH(text):
    P(text, align="l", keep=True, after=2)


def FIG(fn, cap, w=7.9):
    doc.add_picture(os.path.join(IMG, fn), width=Cm(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    P(cap, align="c", size=9, after=8)


def borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge in ("top", "bottom"):
        if kw.get(edge):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), str(kw[edge]))
            e.set(qn("w:color"), "000000")
            b.append(e)
    tcPr.append(b)


def TABLE(cap, rows, widths):
    P(cap, align="c", size=9, after=3, keep=True)
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.width = Cm(widths[j])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            r.bold = i == 0
            borders(c, top=8 if i == 0 else 0, bottom=8 if i in (0, len(rows) - 1) else 0)
    fix_table(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def abstract(label, text, size=10):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(label)
    r.bold = r.italic = True
    r.font.size = Pt(size)
    r2 = p.add_run(" " + text)
    r2.italic = True
    r2.font.size = Pt(size)


# ---------------------------------------------------------- ANGKA TURUNAN
f = m["final"]
cm = m["cm"]
pos = cm["tp"] + cm["fn"]
neg = cm["tn"] + cm["fp"]
by = {r["PAY_1"]: r for r in m["rate_by_pay1"]}
n_ok = sum(by[k]["size"] for k in by if k <= 0)
r_ok = sum(by[k]["size"] * by[k]["mean"] for k in by if k <= 0) / n_ok
numcoef = [c for c in m["coef"] if c["fitur"] in META["numeric"]]
pay1 = next(c for c in m["coef"] if c["fitur"] == "PAY_1")
or_month = math.exp(pay1["koef"] / META["scaler_scale"]["PAY_1"])
thr = {round(r["threshold"], 2): r for r in m["threshold_table"]}
base_m, cw_m, sm_m = m["perbandingan"]
naive_acc = m["n_tidak"] / m["n_akhir"]

# ------------------------------------------------------------ HALAMAN 1
P("PREDIKSI RISIKO GAGAL BAYAR KARTU KREDIT MENGGUNAKAN REGRESI LOGISTIK "
  "DAN SMOTE-NC BERBASIS WEB FLASK", bold=True, align="c", size=15, after=14)
P("Penulis", align="c", size=12, after=2)
P(NAMA, align="c", size=12, after=2)
pe = P("e-mail: ", align="c", size=12, after=14)
r = pe.add_run(EMAIL)
r.font.color.rgb = RGBColor(0x1A, 0x73, 0xE8)
r.font.size = Pt(12)

abs_id = (
    f"Gagal bayar kartu kredit menimbulkan kerugian bagi lembaga keuangan sehingga diperlukan sistem yang dapat "
    f"menyaring nasabah berisiko secara dini. Penelitian ini membangun sistem prediksi risiko gagal bayar dengan "
    f"algoritma Regresi Logistik menggunakan dataset Default of Credit Card Clients yang berisi {m['n_akhir']:,} "
    f"nasabah dengan 23 fitur demografis, riwayat pembayaran, tagihan, dan pembayaran enam bulan. Proporsi nasabah "
    f"gagal bayar hanya {fx(m['rate_default'] * 100, 1)}% sehingga data tidak seimbang; ketidakseimbangan ini "
    f"ditangani dengan SMOTE-NC pada data latih, disertai StandardScaler pada fitur numerik. Pada {m['n_test']:,} "
    f"data uji, model mencapai AUC {fx(f['auc'], 3)}, Recall kelas gagal bayar {fx(f['recall'] * 100, 1)}%, "
    f"Precision {fx(f['precision'] * 100, 1)}%, dan F1-Score {fx(f['f1'] * 100, 1)}%. Analisis menunjukkan "
    f"False Negative (nasabah berisiko yang lolos) lebih merugikan daripada False Positive sehingga Recall "
    f"diprioritaskan, dan tunggakan pembayaran bulan terbaru menjadi prediktor terkuat. Model dan scaler disimpan "
    f"dengan Joblib lalu diintegrasikan ke aplikasi web Flask yang menampilkan keputusan Berisiko atau Tidak "
    f"Berisiko, skor probabilitas, dan faktor pengaruh utama, serta di-deploy ke cloud."
).replace(f"{m['n_akhir']:,}", fx(m["n_akhir"], 0)).replace(f"{m['n_test']:,}", fx(m["n_test"], 0))
abs_en = (
    f"Credit card defaults cause losses for financial institutions, so a system that screens risky customers early is "
    f"needed. This study builds a default-risk prediction system using Logistic Regression on the Default of Credit "
    f"Card Clients dataset containing {m['n_akhir']:,} customers with 23 demographic, repayment-history, bill, and "
    f"payment features. Only {m['rate_default'] * 100:.1f}% of customers default, so the data is imbalanced; this is "
    f"handled with SMOTE-NC on the training set together with StandardScaler on numeric features. On {m['n_test']:,} "
    f"test samples the model achieves an AUC of {f['auc']:.3f}, a default-class Recall of {f['recall'] * 100:.1f}%, "
    f"a Precision of {f['precision'] * 100:.1f}%, and an F1-Score of {f['f1'] * 100:.1f}%. The analysis shows that "
    f"False Negatives (risky customers who slip through) are costlier than False Positives, so Recall is prioritised, "
    f"and the most recent repayment delay is the strongest predictor. The model and scaler are saved with Joblib and "
    f"integrated into a Flask web application that returns a Risky or Not Risky decision, a probability score, and the "
    f"main contributing factors, and is deployed to the cloud."
)
abstract("ABSTRAK:", abs_id)
abstract("Kata kunci:", "regresi logistik, risiko kredit, gagal bayar, SMOTE-NC, Flask")
abstract("ABSTRACT:", abs_en)
abstract("Keywords:", "logistic regression, credit risk, default prediction, SMOTE-NC, Flask")

# ------------------------------------------------- BAGIAN DUA KOLOM
sec2 = doc.add_section(WD_SECTION.NEW_PAGE)
page_setup(sec2)
set_cols(sec2, 2)

H("PENDAHULUAN")
P("Kartu kredit memberi nasabah fleksibilitas pembayaran, tetapi juga membawa risiko bagi penerbit apabila nasabah "
  "gagal membayar tagihan. Kerugian akibat gagal bayar (default) menekan profitabilitas dan memaksa lembaga keuangan "
  "mencadangkan dana lebih besar. Tim manajemen risiko karenanya membutuhkan alat yang mampu memperkirakan "
  "kemungkinan gagal bayar seorang nasabah pada bulan berikutnya berdasarkan data demografis dan riwayat pembayarannya.")
P("Masalah ini merupakan klasifikasi biner (gagal bayar atau tidak). Regresi Linear tidak sesuai karena keluarannya "
  "dapat berada di luar rentang [0, 1] dan tidak dapat dibaca sebagai probabilitas. Regresi Logistik mengatasinya "
  "dengan fungsi sigmoid σ(z) = 1 / (1 + e⁻ᶻ) yang memetakan keluaran linear ke rentang 0 sampai 1, sehingga hasilnya "
  "dapat ditafsirkan sebagai probabilitas kelas positif dan dibandingkan dengan ambang batas keputusan.")
P("Tantangan utama pada data kredit adalah ketidakseimbangan kelas: nasabah gagal bayar jauh lebih sedikit daripada "
  "yang lancar. Pada kondisi ini akurasi menyesatkan, karena model yang selalu menebak \"lancar\" sudah mencapai "
  f"akurasi {fx(naive_acc * 100, 1)}%. Oleh karena itu penelitian ini memakai Confusion Matrix, Precision, Recall, "
  "F1-Score, dan AUC, serta teknik penyeimbangan data.")
P("Tujuan penelitian ini adalah: (1) membangun model Regresi Logistik untuk memprediksi gagal bayar kartu kredit, "
  "(2) menangani class imbalance dan mengevaluasi model secara kritis dengan metrik yang tepat, (3) menginterpretasikan "
  "hasil untuk memberi wawasan manajemen risiko, dan (4) mengimplementasikan model dalam aplikasi web Flask yang "
  "di-deploy sebagai alat pendukung keputusan analis kredit.")

H("METODA")
SH("Dataset")
P(f"Dataset yang digunakan adalah Default of Credit Card Clients (Yeh & Lien, 2009) dari UCI Machine Learning "
  f"Repository yang juga tersedia di Kaggle ({LINK_KAGGLE}). Berkas yang dipakai berisi {fx(m['n_awal'], 0)} "
  f"nasabah. Fitur terdiri atas campuran numerik (LIMIT_BAL, AGE, tagihan BILL_AMT1–6, dan pembayaran PAY_AMT1–6) "
  f"dan kategorikal (SEX, EDUCATION, MARRIAGE), serta status pembayaran PAY_1–6. Target biner adalah gagal bayar "
  f"pada bulan berikutnya (1 = gagal bayar). Seluruh nilai uang dalam dolar Taiwan baru (NT$).")
SH("Pra-pemrosesan Data")
P(f"Data diperiksa untuk nilai hilang ({m['n_missing']} ditemukan) dan baris duplikat ({m['n_dup']} dihapus), "
  f"menyisakan {fx(m['n_akhir'], 0)} baris. Kode kategori yang tidak terdokumentasi dirapikan: EDUCATION bernilai "
  f"0, 5, dan 6 digabung menjadi kategori lainnya, dan MARRIAGE bernilai 0 digabung ke kategori lainnya. Variabel "
  f"kategorikal dikonversi menjadi dummy (baseline: laki-laki, pascasarjana, menikah). Kode status pembayaran "
  f"PAY_1–6 bernilai −2, −1, dan 0 sama-sama berarti tidak ada tunggakan, sehingga diubah menjadi lama tunggakan "
  f"dalam bulan (0 sampai 9) agar hubungannya dengan risiko monoton dan sesuai untuk model linear.")
SH("Pembagian Data, Scaling, dan SMOTE-NC")
P(f"Data dibagi 80:20 secara stratified menjadi {fx(m['n_train'], 0)} data latih dan {fx(m['n_test'], 0)} data uji "
  f"(random_state = 42). StandardScaler di-fit hanya pada fitur numerik data latih lalu diterapkan ke data uji "
  f"untuk mencegah kebocoran data. Ketidakseimbangan kelas pada data latih ({fx(m['train_default_before'], 0)} "
  f"gagal bayar berbanding {fx(m['train_tidak_before'], 0)} lancar) ditangani dengan SMOTE-NC yang membangkitkan "
  f"sampel sintetis kelas minoritas sekaligus menjaga kolom dummy tetap bernilai 0 atau 1. SMOTE-NC hanya "
  f"diterapkan pada data latih sehingga menghasilkan {fx(m['n_train_smote'], 0)} sampel seimbang, sedangkan data uji "
  f"dibiarkan dalam distribusi aslinya.")
SH("Model dan Evaluasi")
P("Model Regresi Logistik (LogisticRegression, solver lbfgs) menghitung z = β₀ + β₁x₁ + … + βₖxₖ lalu "
  "p = σ(z). Prediksi kelas 1 diberikan bila p ≥ 0,5. Sebagai pembanding dilatih model tanpa penyeimbangan dan "
  "model dengan class_weight = balanced. Evaluasi memakai Confusion Matrix, Precision = TP/(TP+FP), "
  "Recall = TP/(TP+FN), F1-Score, dan AUC dari kurva ROC. Model dan scaler disimpan dengan Joblib.")
SH("Aplikasi Web dan Deployment")
P("Aplikasi dibangun dengan Flask. Formulir menerima batas kredit, usia, jenis kelamin, pendidikan, status "
  "pernikahan, serta riwayat 6 bulan (tunggakan, tagihan, pembayaran). Backend memvalidasi input, menerapkan "
  "pra-pemrosesan dan scaler yang sama dengan saat pelatihan, lalu mengembalikan keputusan Berisiko atau Tidak "
  "Berisiko, skor probabilitas, tingkat risiko, dan lima faktor dengan kontribusi terbesar (koefisien dikalikan "
  "nilai terstandarisasi). Antarmuka memakai HTML, CSS, dan JavaScript responsif dengan mode terang dan gelap. "
  f"Aplikasi di-deploy ke cloud dengan gunicorn dan dapat diakses di {LINK_APP}. Kode sumber: {LINK_GITHUB}. "
  f"Video presentasi: {LINK_YOUTUBE}.")

H("HASIL DAN PEMBAHASAN")
SH("Eksplorasi Data dan Class Imbalance")
FIG("eda_target.png", "Gambar 1. Distribusi variabel target", w=6.6)
P(f"Dari {fx(m['n_akhir'], 0)} nasabah, {fx(m['n_default'], 0)} ({fx(m['rate_default'] * 100, 1)}%) gagal bayar dan "
  f"{fx(m['n_tidak'], 0)} lancar (Gambar 1). Rasio sekitar 1 : {fx(m['n_tidak'] / m['n_default'], 1)} ini menegaskan "
  f"perlunya penanganan imbalance dan metrik selain akurasi.")
FIG("eda_pay1.png", "Gambar 2. Tingkat gagal bayar menurut status pembayaran terakhir", w=7.9)
P(f"Status pembayaran terakhir adalah sinyal paling tajam. Nasabah tanpa tunggakan gagal bayar sebesar "
  f"{fx(r_ok * 100, 1)}%, naik menjadi {fx(by[1]['mean'] * 100, 1)}% bila menunggak 1 bulan dan "
  f"{fx(by[2]['mean'] * 100, 1)}% bila menunggak 2 bulan (Gambar 2). Pola ini mendasari pengubahan kode PAY "
  f"menjadi lama tunggakan.")

SH("Perbandingan Strategi Penyeimbangan")
TABLE("Tabel 1. Perbandingan strategi pada data uji", [
    ["Strategi", "Prec.", "Recall", "F1", "AUC"],
    ["Tanpa penyeimbangan", fx(base_m["precision"]), fx(base_m["recall"]), fx(base_m["f1"]), fx(base_m["auc"], 3)],
    ["class_weight balanced", fx(cw_m["precision"]), fx(cw_m["recall"]), fx(cw_m["f1"]), fx(cw_m["auc"], 3)],
    ["SMOTE-NC (final)", fx(sm_m["precision"]), fx(sm_m["recall"]), fx(sm_m["f1"]), fx(sm_m["auc"], 3)],
], [3.6, 1.2, 1.3, 1.1, 1.3])
P(f"Tanpa penyeimbangan, model memiliki Precision tinggi ({fx(base_m['precision'])}) tetapi Recall hanya "
  f"{fx(base_m['recall'])}, artinya sebagian besar nasabah gagal bayar tidak terdeteksi. SMOTE-NC menaikkan "
  f"Recall menjadi {fx(sm_m['recall'])} dengan Precision turun menjadi {fx(sm_m['precision'])}. Hasil "
  f"class_weight hampir sama dengan SMOTE-NC; SMOTE-NC dipilih sesuai ketentuan tugas dan karena "
  f"membuat data latih seimbang secara eksplisit. AUC ketiganya relatif sama ({fx(min(base_m['auc'], cw_m['auc'], sm_m['auc']), 3)}–"
  f"{fx(max(base_m['auc'], cw_m['auc'], sm_m['auc']), 3)}), yang wajar karena penyeimbangan terutama menggeser ambang keputusan, bukan kemampuan "
  f"model membedakan kelas.")

SH("Confusion Matrix")
FIG("model_cm.png", "Gambar 3. Confusion Matrix pada data uji", w=6.4)
P(f"Dari {fx(pos, 0)} nasabah gagal bayar pada data uji, {fx(cm['tp'], 0)} terdeteksi (True Positive) dan "
  f"{fx(cm['fn'], 0)} terlewat (False Negative). Dari {fx(neg, 0)} nasabah lancar, {fx(cm['tn'], 0)} dinilai benar "
  f"(True Negative) dan {fx(cm['fp'], 0)} salah ditandai berisiko (False Positive). False Positive berarti "
  f"nasabah baik harus ditinjau ulang atau limitnya ditahan sehingga pendapatan bunga dan kepuasan nasabah "
  f"berkurang. False Negative berarti nasabah yang akan gagal bayar lolos dan menimbulkan kredit macet. "
  f"Kerugian pokok dari False Negative umumnya jauh lebih besar daripada biaya meninjau nasabah yang sebenarnya "
  f"baik, sehingga False Negative lebih merugikan perusahaan.")

SH("Precision, Recall, F1-Score, dan AUC")
TABLE("Tabel 2. Metrik evaluasi model final", [
    ["Metrik", "Nilai"],
    ["Accuracy", fx(f["accuracy"], 4)],
    ["Precision (gagal bayar)", fx(f["precision"], 4)],
    ["Recall (gagal bayar)", fx(f["recall"], 4)],
    ["F1-Score (gagal bayar)", fx(f["f1"], 4)],
    ["AUC", fx(f["auc"], 4)],
], [4.6, 2.4])
P(f"Dari seluruh nasabah yang ditandai berisiko, {fx(f['precision'] * 100, 1)}% benar-benar gagal bayar "
  f"(Precision). Model berhasil menangkap {fx(f['recall'] * 100, 1)}% dari seluruh nasabah gagal bayar (Recall). "
  f"Untuk manajemen risiko, Recall kelas gagal bayar adalah metrik yang paling penting karena melewatkan nasabah "
  f"berisiko lebih mahal daripada salah menandai nasabah baik; Precision dan F1-Score menjaga agar beban peninjauan "
  f"manual tetap terkendali. Akurasi {fx(f['accuracy'] * 100, 1)}% tidak dijadikan acuan utama karena bias oleh "
  f"kelas mayoritas.")
FIG("model_roc.png", "Gambar 4. Kurva ROC", w=6.4)
P(f"Kurva ROC berada di atas garis tebakan acak dengan AUC {fx(f['auc'], 3)}. Nilai ini tergolong baik (rentang "
  f"0,7–0,8): bila diambil satu nasabah gagal bayar dan satu nasabah lancar secara acak, model memberi skor lebih "
  f"tinggi kepada nasabah gagal bayar pada sekitar {fx(f['auc'] * 100, 0)}% pasangan. Nilai yang jauh dari 1,0 juga "
  f"mengindikasikan tidak adanya kebocoran data.")

SH("Trade-off Ambang Batas")
TABLE("Tabel 3. Pengaruh ambang batas (data uji)", [
    ["Threshold", "Prec.", "Recall", "FP", "FN"],
] + [[fx(t, 2), fx(thr[t]["precision"]), fx(thr[t]["recall"]), fx(thr[t]["fp"], 0), fx(thr[t]["fn"], 0)]
     for t in (0.3, 0.4, 0.5, 0.6)], [2.0, 1.4, 1.5, 1.4, 1.4])
P(f"Menurunkan threshold dari 0,50 ke 0,40 menaikkan Recall dari {fx(thr[0.5]['recall'])} menjadi "
  f"{fx(thr[0.4]['recall'])} dan menurunkan False Negative dari {fx(thr[0.5]['fn'], 0)} menjadi "
  f"{fx(thr[0.4]['fn'], 0)}, tetapi False Positive naik dari {fx(thr[0.5]['fp'], 0)} menjadi "
  f"{fx(thr[0.4]['fp'], 0)} dan Precision turun ke {fx(thr[0.4]['precision'])}. Ambang 0,50 dipakai pada "
  f"aplikasi sesuai ketentuan modul, namun lembaga yang lebih menghindari risiko dapat menurunkannya dengan "
  f"konsekuensi beban peninjauan manual yang lebih besar.")

SH("Interpretasi Koefisien")
FIG("model_koefisien.png", "Gambar 5. Koefisien numerik terbesar", w=7.9)
P(f"Di antara fitur numerik terstandarisasi, tunggakan bulan ke-1 memiliki koefisien positif terbesar "
  f"({fx(pay1['koef'], 3)}; odds ratio {fx(pay1['odds_ratio'], 2)} per kenaikan satu simpangan baku). Dalam satuan "
  f"asli, setiap tambahan satu bulan tunggakan terbaru menaikkan odds gagal bayar sekitar {fx(or_month, 1)} kali. "
  f"Pembayaran bulan ke-2 ({fx(numcoef[1]['koef'], 3)}) dan batas kredit ({fx(numcoef[2]['koef'], 3)}) berkoefisien "
  f"negatif, artinya pembayaran yang lebih besar dan limit yang lebih tinggi dikaitkan dengan risiko lebih rendah. "
  f"Koefisien tagihan dan pembayaran antarbulan saling berkorelasi sehingga tidak boleh ditafsirkan terpisah. "
  f"Koefisien kategori langka (misalnya pendidikan lainnya) juga kurang stabil karena sampelnya sedikit.")

SH("Aplikasi Web")
doc.add_picture(os.path.join("assets", "screenshot_app.png"), width=Cm(8.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
P("Gambar 6. Tampilan aplikasi dan hasil prediksi", align="c", size=9, after=8)
P("Aplikasi menerima data nasabah baru dan menampilkan prediksi seketika berupa keputusan Berisiko atau Tidak "
  "Berisiko, skor probabilitas pada gauge, tingkat risiko (rendah, sedang, tinggi), saran tindak lanjut, dan faktor "
  "utama yang menaikkan atau menurunkan risiko dibanding nasabah rata-rata. Karena model dilatih dengan data yang "
  "diseimbangkan, skor probabilitasnya cenderung lebih tinggi daripada frekuensi gagal bayar sebenarnya; "
  "skor sebaiknya dibaca sebagai peringkat risiko relatif.")

SH("Rekomendasi Bisnis dan Keterbatasan")
P("Bagi tim manajemen risiko, nasabah yang mulai menunggak perlu menjadi prioritas pemantauan dan penagihan dini "
  "karena tunggakan adalah prediktor terkuat. Nasabah dengan skor tinggi dapat diarahkan ke peninjauan manual, "
  "pembatasan kenaikan limit, atau program restrukturisasi, sedangkan nasabah berskor rendah dapat diproses lebih "
  "cepat. Keterbatasan penelitian ini: data berasal dari satu bank di Taiwan tahun 2005 sehingga perlu validasi "
  "ulang pada data terkini, model linear tidak menangkap interaksi antarfitur, ambang keputusan belum dikaitkan "
  "dengan biaya finansial nyata, dan penggunaan atribut demografis seperti jenis kelamin dalam keputusan kredit "
  "perlu ditinjau dari sisi kepatuhan dan keadilan.")

H("SIMPULAN DAN SARAN")
P(f"Penelitian ini berhasil membangun sistem prediksi risiko gagal bayar kartu kredit dengan Regresi Logistik, "
  f"StandardScaler, dan SMOTE-NC. Pada data uji model mencapai AUC {fx(f['auc'], 3)}, Recall "
  f"{fx(f['recall'] * 100, 1)}%, Precision {fx(f['precision'] * 100, 1)}%, dan F1-Score "
  f"{fx(f['f1'] * 100, 1)}%. Penanganan class imbalance menaikkan Recall dari {fx(base_m['recall'] * 100, 1)}% "
  f"menjadi {fx(sm_m['recall'] * 100, 1)}%. Tunggakan pembayaran terbaru terbukti sebagai prediktor terkuat, dan "
  f"False Negative merupakan kesalahan yang paling merugikan sehingga Recall diprioritaskan. Model berhasil "
  f"dikemas dalam aplikasi web Flask yang di-deploy dan dapat dipakai analis kredit sebagai alat pendukung "
  f"keputusan.")
P("Saran pengembangan: membandingkan dengan model non-linear seperti Random Forest atau Gradient Boosting, "
  "mengkalibrasi probabilitas dan menentukan threshold berdasarkan biaya finansial, serta menguji model pada data "
  "yang lebih baru dan melakukan audit keadilan atas atribut demografis.")

H("DAFTAR PUSTAKA")
for ref in [
    "Yeh, I. C., & Lien, C. H. (2009). The comparisons of data mining techniques for the predictive accuracy of probability of default of credit card clients. Expert Systems with Applications, 36(2), 2473–2480.",
    f"UCI Machine Learning Repository / Kaggle. (n.d.). Default of Credit Card Clients Dataset. {LINK_KAGGLE}",
    "Chawla, N. V., Bowyer, K. W., Hall, L. O., & Kegelmeyer, W. P. (2002). SMOTE: Synthetic Minority Over-sampling Technique. Journal of Artificial Intelligence Research, 16, 321–357.",
    "Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    "Lemaître, G., Nogueira, F., & Aridas, C. K. (2017). Imbalanced-learn: A Python toolbox to tackle the curse of imbalanced datasets in machine learning. Journal of Machine Learning Research, 18(17), 1–5.",
    "James, G., Witten, D., Hastie, T., & Tibshirani, R. (2013). An Introduction to Statistical Learning. Springer.",
    "Pallets Projects. (n.d.). Flask Documentation. https://flask.palletsprojects.com/",
]:
    P(ref, align="j", after=4)

doc.save(OUT)
print("Tersimpan:", OUT)
